#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script tự động hóa tạo SHEET TỰ CHẤM ĐIỂM & RÀ SOÁT CHẤT LƯỢNG HỒ SƠ KIỂM TOÁN
theo đúng chuẩn Bảng chấm điểm của Bộ Tài chính & VACPA 2014.
Áp dụng cho 13 file Giấy làm việc (GLV) mẫu trong thư mục GLV MAU.
"""

import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

TARGET_DIR = 'D:/Desktop/Project/5. AuditSoft/GLV MAU'

# Cơ sở dữ liệu 13 bảng chấm điểm theo từng phần hành
SCORING_SHEETS_DATABASE = {
    'D100': {
        'sheet_name': 'D 199_ChamDiem',
        'ref_code': 'D 199',
        'section_title': 'PHẦN HÀNH TIỀN VÀ CÁC KHOẢN TƯƠNG ĐƯƠNG TIỀN (VSA 330, 500, 505)',
        'max_points': 6.0,
        'criteria': [
            ('1.1', 'KTV có tham gia chứng kiến kiểm kê quỹ tiền mặt tồn quỹ không? (Trường hợp kiểm kê trước/sau ngày kết thúc kỳ, có thực hiện đối chiếu biến động đến ngày 31/12 không?)', 'VSA 330, VSA 500', 1.0, 'D 146', 'Đã kẹp Biên bản kiểm kê quỹ tiền mặt'),
            ('1.2', 'Có thực hiện thủ tục đối chiếu số liệu tiền mặt thực tế kiểm kê và số liệu sổ kế toán, BCTC không? Chênh lệch (nếu có) đã được giải thích hoặc điều chỉnh hợp lý chưa?', 'VSA 330, VSA 500', 1.0, 'D 146, D198', 'Khớp số liệu thực tế với Sổ cái'),
            ('1.3', 'Có thực hiện thủ tục gửi thư xác nhận (TXN) số dư tài khoản tiền gửi ngân hàng (100% tài khoản ngân hàng mở trong kỳ) không?', 'VSA 330, VSA 505', 1.0, 'D 141', 'Đã gửi thư xác nhận 100% ngân hàng'),
            ('1.4', 'Có thực hiện thủ tục đối chiếu số liệu thư xác nhận số dư TGNH (hoặc biên bản đối chiếu số dư / sổ phụ ngân hàng) với số liệu Sổ cái kế toán, BCTC không?', 'VSA 330, VSA 505', 1.0, 'D 143, D 110', 'Khớp số dư thư xác nhận và sổ phụ NH'),
            ('1.5', 'Có thực hiện các thủ tục kiểm toán cần thiết khác (chọn mẫu thu chi, kiểm tra Cut-off trước và sau 31/12...) để đảm bảo thu thập đầy đủ bằng chứng thích hợp không?', 'VSA 330, VSA 530', 1.0, 'D 191.1, D 195TM, D 196', 'Đã kiểm tra chọn mẫu chứng từ và Cut-off'),
            ('1.6', 'Có thực hiện thủ tục kiểm tra việc phân loại và trình bày số liệu tiền trên Báo cáo lưu chuyển tiền tệ và Thuyết minh BCTC không?', 'VSA 330', 1.0, 'D 110.1', 'Trình bày phù hợp chuẩn mực kế toán'),
        ]
    },
    'D200': {
        'sheet_name': 'D 299_ChamDiem',
        'ref_code': 'D 299',
        'section_title': 'PHẦN HÀNH CÁC KHOẢN ĐẦU TƯ TÀI CHÍNH (VSA 330, 500, 505, 540)',
        'max_points': 5.0,
        'criteria': [
            ('2.1', 'Có xem xét tính hiện hữu hoặc kiểm tra các giấy tờ cần thiết chứng minh quyền sở hữu của đơn vị đối với các khoản đầu tư tài chính tại ngày kết thúc kỳ kế toán không?', 'VSA 330, VSA 500', 1.0, 'D240', 'Đầy đủ hồ sơ pháp lý & chứng nhận sở hữu'),
            ('2.2', 'Có chọn mẫu và gửi thư xác nhận (TXN) số dư các khoản đầu tư với bên thứ ba? Có đối chiếu số liệu TXN và sổ kế toán, BCTC? Chênh lệch đã được giải thích/điều chỉnh chưa?', 'VSA 330, VSA 505', 1.0, 'D241', 'Đã gửi và đối chiếu số dư xác nhận'),
            ('2.3', 'Có thực hiện thủ tục kiểm toán thay thế đối với các khoản đầu tư có số dư lớn nhưng không thể gửi TXN hoặc không nhận được phản hồi TXN không?', 'VSA 505', 1.0, 'D241, D252', 'Đã thực hiện thủ tục kiểm tra thay thế'),
            ('2.4', 'Có kiểm tra việc trích lập dự phòng giảm giá các khoản đầu tư tài chính theo đúng quy định hiện hành không?', 'VSA 330, VSA 540', 1.0, 'D252, D298', 'Trích lập dự phòng đúng chuẩn mực'),
            ('2.5', 'Có thực hiện các thủ tục kiểm toán cần thiết khác để đảm bảo thu thập đầy đủ bằng chứng kiểm toán thích hợp về khoản mục đầu tư tài chính không?', 'VSA 330', 1.0, 'D210, D230', 'Đạt yêu cầu thu thập bằng chứng'),
        ]
    },
    'D300': {
        'sheet_name': 'D 399_ChamDiem',
        'ref_code': 'D 399',
        'section_title': 'PHẦN HÀNH CÁC KHOẢN PHẢI THU KHÁCH HÀNG (VSA 330, 500, 505, 240, 540)',
        'max_points': 6.0,
        'criteria': [
            ('3.1', 'Có chọn mẫu và gửi thư xác nhận (TXN) số dư các khoản phải thu với bên thứ ba (các khách hàng có số dư trọng yếu) không?', 'VSA 330, VSA 505', 1.0, 'D 340, D 341', 'Đã gửi TXN theo mẫu chọn'),
            ('3.2', 'Có thực hiện thủ tục đối chiếu số liệu TXN / Biên bản đối chiếu công nợ với số liệu sổ kế toán, BCTC? Chênh lệch (nếu có) đã được giải thích/điều chỉnh hợp lý chưa?', 'VSA 330, VSA 505', 1.0, 'D 341, D 398', 'Chênh lệch đối chiếu đã được giải trình rõ'),
            ('3.3', 'Có thực hiện thủ tục kiểm toán thay thế đối với các khoản phải thu lớn không nhận được phản hồi TXN (kiểm tra hóa đơn, vận đơn, chứng từ thu tiền sau kỳ)?', 'VSA 505', 1.0, 'D 351.1, D 351.2', 'Đã kiểm tra chứng từ thu tiền sau ngày 31/12'),
            ('3.4', 'Có thực hiện thủ tục kiểm toán đối với các khoản khách hàng trả tiền trước ngắn hạn/dài hạn không?', 'VSA 330, VSA 240', 1.0, 'D353', 'Khách hàng trả trước có căn cứ hợp đồng'),
            ('3.5', 'Có kiểm tra việc phân tích tuổi nợ và trích lập dự phòng nợ phải thu khó đòi theo quy định hiện hành không?', 'VSA 330, VSA 540', 1.0, 'D 352', 'Dự phòng trích lập đúng quy định'),
            ('3.6', 'Có thực hiện các thủ tục kiểm toán cần thiết khác (chọn mẫu nghiệp vụ, kiểm tra Cut-off doanh thu/phải thu trước và sau 31/12)?', 'VSA 330, VSA 560', 1.0, 'D 310, D354, D 390', 'Đã kiểm tra Cut-off và mẫu nghiệp vụ'),
        ]
    },
    'D500': {
        'sheet_name': 'D 599_ChamDiem',
        'ref_code': 'D 599',
        'section_title': 'PHẦN HÀNH HÀNG TỒN KHO (VSA 501, 330, 500, 315, 540)',
        'max_points': 6.0,
        'criteria': [
            ('4.1', 'Kiểm toán viên có tham gia kiểm kê hiện vật hàng tồn kho không? (Đối chiếu, kiểm tra biến động số dư hàng tồn kho nếu kiểm kê trước/sau ngày kết thúc kỳ)?', 'VSA 501, VSA 330', 1.0, 'D540, D541', 'Đã tham gia chứng kiến kiểm kê kho'),
            ('4.2', 'Có thực hiện thủ tục đối chiếu số liệu thực tế kiểm kê và số liệu sổ kế toán, thẻ kho? Chênh lệch (nếu có) đã được giải thích hoặc điều chỉnh hợp lý chưa?', 'VSA 501, VSA 500', 1.0, 'D540, D598', 'Khớp số liệu thẻ kho và sổ kế toán'),
            ('4.3', 'Có kiểm tra phương pháp tính giá hàng tồn kho không (tính phù hợp và tính nhất quán so với năm trước)?', 'VSA 330, VSA 315', 1.0, 'D553, D554', 'Áp dụng phương pháp tính giá nhất quán'),
            ('4.4', 'Có kiểm tra việc tập hợp chi phí sản xuất, tính giá thành và xác định giá trị sản phẩm dở dang cuối kỳ không?', 'VSA 330', 1.0, 'D556, D556.1', 'Phương pháp tính giá thành phù hợp'),
            ('4.5', 'Có kiểm tra việc lập dự phòng giảm giá hàng tồn kho (so sánh giá gốc với giá trị thuần có thể thực hiện được) không?', 'VSA 330, VSA 540', 1.0, 'D557', 'Dự phòng giảm giá trích lập đầy đủ'),
            ('4.6', 'Có thực hiện các thủ tục kiểm toán cần thiết khác (kiểm tra Cut-off nhập xuất kho, hàng đi đường, hàng gửi bán)?', 'VSA 330, VSA 560', 1.0, 'D550, D558, D595', 'Đã kiểm tra Cut-off và hàng gửi bán'),
        ]
    },
    'D600': {
        'sheet_name': 'D 699_ChamDiem',
        'ref_code': 'D 699',
        'section_title': 'PHẦN HÀNH CHI PHÍ TRẢ TRƯỚC (VSA 330, 530, 315, 540)',
        'max_points': 3.0,
        'criteria': [
            ('5.1', 'Có kiểm tra chọn mẫu đối với việc ghi nhận các khoản chi phí trả trước phát sinh trong kỳ không?', 'VSA 330, VSA 530', 1.0, 'D 640, D 641', 'Hóa đơn chứng từ phát sinh hợp lệ'),
            ('5.2', 'Có kiểm tra tiêu thức và thời gian phân bổ đối với các khoản chi phí trả trước không (tính phù hợp, tính nhất quán)?', 'VSA 330, VSA 540', 1.0, 'D 640', 'Tiêu thức phân bổ nhất quán'),
            ('5.3', 'Có thực hiện các thủ tục kiểm toán cần thiết khác (đối chiếu Sổ cái, kiểm tra phân loại ngắn hạn và dài hạn trên BCTC)?', 'VSA 330', 1.0, 'D 610, D 610.1', 'Phân loại ngắn/dài hạn chính xác'),
        ]
    },
    'D700': {
        'sheet_name': 'D 799_ChamDiem',
        'ref_code': 'D 799',
        'section_title': 'PHẦN HÀNH TÀI SẢN CỐ ĐỊNH & XÂY DỰNG CƠ BẢN DỞ DANG (VSA 330, 500, 540, 520)',
        'max_points': 6.0,
        'criteria': [
            ('6.1', 'Có kiểm tra chọn mẫu đối với các TSCĐ hữu hình, vô hình, thuê tài chính và XDCBDD tăng, giảm trong kỳ (hóa đơn, hợp đồng, nghiệm thu)?', 'VSA 330, VSA 530', 1.0, 'D 740, D 741', 'Chứng từ tăng giảm TSCĐ đầy đủ'),
            ('6.2', 'Có kiểm tra tính hiện hữu hoặc tình trạng hoạt động thực tế của các tài sản cố định trọng yếu không?', 'VSA 330, VSA 500', 1.0, 'D 740', 'Đã kiểm tra hiện hữu tài sản'),
            ('6.3', 'Có kiểm tra Bảng tính, phương pháp và khung thời gian tính khấu hao TSCĐ theo quy định hiện hành (TT 45/2013/TT-BTC)?', 'VSA 330, VSA 540', 1.0, 'D 790', 'Khung thời gian khấu hao đúng quy định'),
            ('6.4', 'Có thực hiện ước tính độc lập về chi phí khấu hao trong kỳ để so sánh với số liệu của đơn vị được kiểm toán không?', 'VSA 520, VSA 540', 1.0, 'D 792', 'Chênh lệch ước tính không trọng yếu'),
            ('6.5', 'Có kiểm tra tính đầy đủ, phù hợp của thuyết minh liên quan đến TSCĐ, chi phí XDCBDD trên BCTC (thế chấp, bảo lãnh, mục đích sử dụng)?', 'VSA 330', 1.0, 'D720', 'Thuyết minh đầy đủ theo chuẩn mực'),
            ('6.6', 'Có kiểm tra tiến độ, nghiệm thu và điều kiện kết chuyển chi phí xây dựng cơ bản dở dang không?', 'VSA 330', 1.0, 'D794', 'Kết chuyển XDCBDD đúng thời điểm bàn giao'),
        ]
    },
    'E100': {
        'sheet_name': 'E 199_ChamDiem',
        'ref_code': 'E 199',
        'section_title': 'PHẦN HÀNH VAY VÀ NỢ THUÊ TÀI CHÍNH (VSA 330, 505, 520, 540)',
        'max_points': 6.0,
        'criteria': [
            ('7.1', 'Có gửi thư xác nhận (TXN) số dư các khoản vay và nợ với bên thứ ba (100% ngân hàng và tổ chức tín dụng mở trong kỳ) không?', 'VSA 330, VSA 505', 1.0, 'E 140, E141', 'Đã gửi xác nhận ngân hàng đầy đủ'),
            ('7.2', 'Có thực hiện thủ tục đối chiếu số liệu TXN / Biên bản đối chiếu nợ vay với số liệu Sổ cái kế toán và BCTC? Chênh lệch đã được giải thích chưa?', 'VSA 330, VSA 505', 1.0, 'E 141', 'Khớp số dư nợ gốc và lãi vay'),
            ('7.3', 'Có thực hiện thủ tục kiểm toán thay thế đối với các khoản vay không thể gửi hoặc không nhận được phản hồi TXN?', 'VSA 505', 1.0, 'E 150', 'Kiểm tra khế ước và chứng từ trả nợ'),
            ('7.4', 'Có kiểm tra và thực hiện ước tính độc lập chi phí lãi vay trong kỳ và so sánh với số liệu của đơn vị không?', 'VSA 520, VSA 540', 1.0, 'E 152', 'Chi phí lãi vay khớp ước tính'),
            ('7.5', 'Có kiểm tra việc phân loại các khoản nợ vay ngắn hạn và dài hạn đến hạn trả trong vòng 12 tháng không?', 'VSA 330', 1.0, 'E 110.1', 'Phân loại nợ đến hạn đúng chuẩn'),
            ('7.6', 'Có kiểm tra tính đầy đủ của thuyết minh các khoản vay (tài sản bảo đảm, cam kết tài chính, vi phạm hợp đồng tín dụng)?', 'VSA 330', 1.0, 'E 120', 'Thuyết minh nợ vay đầy đủ'),
        ]
    },
    'E200': {
        'sheet_name': 'E 299_ChamDiem',
        'ref_code': 'E 299',
        'section_title': 'PHẦN HÀNH PHẢI TRẢ NGƯỜI BÁN & CHI PHÍ TRÍCH TRƯỚC (VSA 330, 505, 560, 540)',
        'max_points': 9.0,
        'criteria': [
            ('8.1', 'Có gửi thư xác nhận (TXN) số dư các khoản phải trả với bên thứ ba (các nhà cung cấp có số dư hoặc phát sinh trọng yếu)?', 'VSA 330, VSA 505', 1.0, 'E240, E241', 'Đã gửi TXN nhà cung cấp'),
            ('8.2', 'Có thực hiện thủ tục đối chiếu số liệu TXN / Biên bản đối chiếu nợ với sổ sách kế toán, BCTC? Chênh lệch đã giải thích chưa?', 'VSA 330, VSA 505', 1.0, 'E241', 'Khớp số dư đối chiếu nợ phải trả'),
            ('8.3', 'Có thực hiện thủ tục kiểm toán thay thế đối với các khoản nợ lớn không nhận được phản hồi TXN?', 'VSA 505', 1.0, 'E242', 'Kiểm tra phiếu nhập kho và thanh toán sau kỳ'),
            ('8.4', 'Có kiểm tra chứng từ các nghiệp vụ kinh tế phát sinh sau ngày 31/12 để tìm kiếm nợ chưa ghi sổ (Unrecorded liabilities - Cutoff)?', 'VSA 330, VSA 560', 1.0, 'E252, E253', 'Không phát hiện nợ chưa ghi sổ'),
            ('8.5', 'Có kiểm tra việc phân loại các khoản nợ phải trả ngắn hạn và dài hạn trên BCTC không?', 'VSA 330', 1.0, 'E 210', 'Phân loại thời hạn nợ chính xác'),
            ('8.6', 'Có thực hiện các thủ tục kiểm toán cần thiết khác đối với khoản mục phải trả người bán?', 'VSA 330', 1.0, 'E220', 'Đạt yêu cầu thu thập bằng chứng'),
            ('11.1', 'Có thực hiện kiểm tra chọn mẫu các khoản chi phí phải trả (TK 335) trích trước còn dư tại thời điểm khóa sổ không?', 'VSA 330, VSA 540', 1.0, 'E260, E270', 'Chi phí trích trước có đủ căn cứ'),
            ('11.2', 'Có kiểm tra các nghiệp vụ phát sinh sau ngày kết thúc kỳ kế toán để xác định chi phí đã phát sinh nhưng chưa được trích trước?', 'VSA 330, VSA 560', 1.0, 'E253', 'Đã rà soát chi phí sau niên độ'),
            ('11.3', 'Có kiểm tra tính hợp lý của việc hoàn nhập hoặc trích thêm chi phí phải trả trong kỳ không?', 'VSA 330, VSA 540', 1.0, 'E260', 'Hạch toán hoàn nhập đúng quy định'),
        ]
    },
    'E300': {
        'sheet_name': 'E 399_ChamDiem',
        'ref_code': 'E 399',
        'section_title': 'PHẦN HÀNH THUẾ VÀ CÁC KHOẢN PHẢI NỘP NHÀ NƯỚC (VSA 250, 330, 520)',
        'max_points': 4.0,
        'criteria': [
            ('9.1', 'Có kiểm tra, đối chiếu các mức thuế suất, cách tính thuế từng sắc thuế với văn bản quy định hiện hành không?', 'VSA 250, VSA 330', 1.0, 'E 341', 'Áp dụng thuế suất đúng luật'),
            ('9.2', 'Có đối chiếu số liệu hạch toán với các tờ khai thuế định kỳ (GTGT, TNCN, TNDN tạm nộp...), kể cả tờ khai các tháng sau 31/12?', 'VSA 330, VSA 250', 1.0, 'E 330, E 340', 'Khớp số liệu tờ khai thuế'),
            ('9.3', 'Có ước tính độc lập số thuế phải nộp (đặc biệt là thuế TNDN hiện hành và rà soát chi phí không được trừ)?', 'VSA 330, VSA 520', 1.0, 'E 382', 'Đã bóc tách chi phí không được trừ'),
            ('9.4', 'Có đối chiếu số tiền thuế đã nộp với Giấy nộp tiền vào NSNN và Thông báo đối chiếu nghĩa vụ của Cơ quan Thuế?', 'VSA 330, VSA 500', 1.0, 'E 380, E 381', 'Đã đối chiếu giấy nộp tiền vào NSNN'),
        ]
    },
    'E400': {
        'sheet_name': 'E 499_ChamDiem',
        'ref_code': 'E 499',
        'section_title': 'PHẦN HÀNH PHẢI TRẢ NGƯỜI LAO ĐỘNG & BẢO HIỂM (VSA 330, 520, 250, 550)',
        'max_points': 5.0,
        'criteria': [
            ('10.1', 'Có thực hiện thủ tục phân tích biến động chi phí tiền lương theo tháng/quý hoặc ước tính độc lập quỹ lương trong kỳ?', 'VSA 330, VSA 520', 1.0, 'E 440', 'Chi phí lương biến động hợp lý'),
            ('10.2', 'Có kiểm tra chọn mẫu hồ sơ bảng lương, đối chiếu với hợp đồng lao động, bảng chấm công và chứng từ chi trả tiền lương?', 'VSA 330, VSA 530', 1.0, 'E 441', 'Bảng lương có chữ ký nhận đầy đủ'),
            ('10.3', 'Có kiểm tra việc phân bổ chi phí tiền lương và các khoản trích theo lương vào giá thành và chi phí SXKD?', 'VSA 330', 1.0, 'E 490', 'Phân bổ đúng đối tượng chi phí'),
            ('10.4', 'Có kiểm tra, đối chiếu các khoản trích nộp BHXH, BHYT, BHTN với Thông báo kết quả đóng BHXH của cơ quan BHXH (C12-TS)?', 'VSA 330, VSA 250', 1.0, 'E 491', 'Khớp đúng số liệu cơ quan BHXH'),
            ('10.5', 'Có kiểm tra và đánh giá tính đầy đủ của thuyết minh về thu nhập của Hội đồng quản trị, Ban Giám đốc trên BCTC?', 'VSA 550', 1.0, 'E 410.1, E 420', 'Thuyết minh thu nhập minh bạch'),
        ]
    },
    'F100': {
        'sheet_name': 'F 199_ChamDiem',
        'ref_code': 'F 199',
        'section_title': 'PHẦN HÀNH NGUỒN VỐN CHỦ SỞ HỮU (VSA 330, 505, 250)',
        'max_points': 6.0,
        'criteria': [
            ('12.1', 'Có kiểm tra chứng từ, tài liệu pháp lý liên quan đối với các nghiệp vụ tăng, giảm vốn chủ sở hữu trong năm?', 'VSA 330, VSA 530', 1.0, 'F141', 'Chứng từ tăng giảm vốn hợp lệ'),
            ('12.2', 'Có đối chiếu Giấy phép ĐKKD, Điều lệ công ty và danh sách cổ đông / thành viên góp vốn tại thời điểm gần nhất?', 'VSA 330, VSA 505', 1.0, 'F140', 'Vốn góp đúng tỷ lệ ĐKKD'),
            ('12.3', 'Có kiểm tra, đối chiếu lợi nhuận sau thuế phát sinh trong năm trên Bảng CĐKT với Báo cáo kết quả HĐKD?', 'VSA 330', 1.0, 'F130', 'Khớp số lợi nhuận sau thuế cả năm'),
            ('12.4', 'Có kiểm tra chứng từ và tính tuân thủ pháp luật đối với các nghiệp vụ phân phối lợi nhuận, chia cổ tức?', 'VSA 330, VSA 250', 1.0, 'F148', 'Biên bản họp ĐHĐCĐ/HĐQT đầy đủ'),
            ('12.5', 'Có kiểm tra việc trích lập và sử dụng các quỹ (quỹ ĐTPT, quỹ khen thưởng phúc lợi) theo đúng Điều lệ và quy định?', 'VSA 330, VSA 250', 1.0, 'F 190', 'Trích lập quỹ đúng tỷ lệ Điều lệ'),
            ('12.6', 'Có kiểm tra việc trình bày Báo cáo biến động vốn chủ sở hữu và thông tin thuyết minh nguồn vốn trên BCTC không?', 'VSA 330', 1.0, 'F120', 'Thuyết minh biến động vốn đầy đủ'),
        ]
    },
    'G100': {
        'sheet_name': 'G 199_ChamDiem',
        'ref_code': 'G 199',
        'section_title': 'PHẦN HÀNH DOANH THU BÁN HÀNG & CUNG CẤP DỊCH VỤ (VSA 240, 330, 520, 500, 560)',
        'max_points': 6.0,
        'criteria': [
            ('14.1', 'Có thực hiện thủ tục phân tích sự biến động của doanh thu theo nhóm mặt hàng, theo thời gian, so với năm trước?', 'VSA 240, VSA 520', 1.0, 'G 140', 'Doanh thu biến động phù hợp'),
            ('14.2', 'Có đối chiếu số liệu hạch toán doanh thu với bộ phận bán hàng, kinh doanh, giao hàng hoặc hợp đồng kinh tế?', 'VSA 240, VSA 500', 1.0, 'G 150', 'Khớp phiếu xuất kho giao hàng'),
            ('14.3', 'Có đối chiếu số liệu doanh thu hạch toán với Bảng kê hóa đơn bán ra và Tờ khai thuế GTGT định kỳ?', 'VSA 240, VSA 500', 1.0, 'G 141', 'Khớp doanh thu kê khai thuế GTGT'),
            ('14.4', 'Có kiểm tra chọn mẫu các nghiệp vụ ghi nhận doanh thu (Hóa đơn GTGT, Phiếu xuất kho, Biên bản bàn giao, nghiệm thu)?', 'VSA 330, VSA 530', 1.0, 'G 191.1', 'Chứng từ bán hàng đầy đủ, hợp lệ'),
            ('14.5', 'Có kiểm tra tính đúng kỳ (Cut-off) của doanh thu trước và sau ngày kết thúc kỳ kế toán 31/12 không?', 'VSA 240, VSA 560', 1.0, 'G 194', 'Không có doanh thu ghi nhận sai kỳ'),
            ('14.6', 'Có kiểm tra các khoản giảm trừ doanh thu (chiết khấu, giảm giá, hàng trả lại) và doanh thu bán hàng cho các bên liên quan?', 'VSA 330, VSA 550', 1.0, 'G 151, G 152', 'Giao dịch bên liên quan đúng quy định'),
        ]
    },
    'G200': {
        'sheet_name': 'G 299_ChamDiem',
        'ref_code': 'G 299',
        'section_title': 'PHẦN HÀNH GIÁ VỐN, CHI PHÍ HOẠT ĐỘNG & TÀI CHÍNH (VSA 330, 520, 530, 560)',
        'max_points': 14.0,
        'criteria': [
            ('15.1', 'Có thực hiện thủ tục phân tích biến động giá vốn hàng bán so với doanh thu, tỷ lệ lãi gộp theo loại hình, so năm trước?', 'VSA 330, VSA 520', 1.0, 'G240', 'Tỷ lệ lãi gộp biến động hợp lý'),
            ('15.2', 'Có kiểm tra tính phù hợp giữa việc ghi nhận doanh thu và ghi nhận giá vốn hàng bán (Matching principle)?', 'VSA 330, VSA 520', 1.0, 'G241', 'Giá vốn phù hợp với doanh thu'),
            ('15.3', 'Có kiểm tra việc ghi nhận giá vốn và đối chiếu với số liệu xuất kho trên Báo cáo Nhập-Xuất-Tồn kho?', 'VSA 330', 1.0, 'G242', 'Khớp giá vốn xuất kho'),
            ('15.4', 'Có kiểm tra tính đúng kỳ (Cut-off) của giá vốn hàng bán trước và sau ngày 31/12 không?', 'VSA 330, VSA 560', 1.0, 'G250', 'Giá vốn ghi nhận đúng kỳ'),
            ('15.5', 'Có kiểm tra chọn mẫu các nghiệp vụ giá vốn bất thường (hao hụt vượt định mức, trích lập dự phòng HTK, ngừng SX)?', 'VSA 330, VSA 530', 1.0, 'G260', 'Không có giá vốn bất thường'),
            ('15.6', 'Có đối chiếu tổng giá vốn trên Sổ cái với Báo cáo kết quả hoạt động kinh doanh không?', 'VSA 330', 1.0, 'G210', 'Khớp số liệu tổng hợp'),
            ('16.1', 'Có thực hiện thủ tục phân tích biến động chi phí bán hàng và chi phí QLDN theo khoản mục và thời gian?', 'VSA 330, VSA 520', 1.0, 'G340, G440', 'Chi phí biến động phù hợp quy mô'),
            ('16.2', 'Có đối chiếu chéo các khoản mục chi phí với các phần hành liên quan (lương, khấu hao, trả trước, trích trước)?', 'VSA 330', 1.0, 'G340, G440', 'Khớp chéo các phần hành liên quan'),
            ('16.3', 'Có chọn mẫu kiểm tra hóa đơn, chứng từ chi phí bán hàng và chi phí quản lý phát sinh trong kỳ?', 'VSA 330, VSA 530', 1.0, 'G390, G490', 'Hóa đơn chứng từ hợp lệ, đầy đủ'),
            ('16.4', 'Có đối chiếu tổng chi phí bán hàng và chi phí quản lý với Báo cáo kết quả hoạt động kinh doanh?', 'VSA 330', 1.0, 'G310, G410', 'Khớp số liệu BCTC'),
            ('17.1', 'Có đối chiếu các khoản mục doanh thu, chi phí hoạt động tài chính với các phần hành liên quan (tiền gửi, vay)?', 'VSA 330, VSA 520', 1.0, 'G220', 'Lãi tiền gửi và lãi vay khớp đối chiếu'),
            ('17.2', 'Có kiểm tra chọn mẫu các khoản doanh thu tài chính và chi phí tài chính (chênh lệch tỷ giá, lãi vay)?', 'VSA 330, VSA 530', 1.0, 'G220', 'Hạch toán tỷ giá đúng quy định'),
            ('17.3', 'Có kiểm tra tính đầy đủ, đúng kỳ của doanh thu, chi phí hoạt động tài chính không?', 'VSA 330', 1.0, 'G220', 'Ghi nhận đúng kỳ kế toán'),
            ('17.4', 'Có kiểm tra các khoản thu nhập khác (TK 711) và chi phí khác (TK 811), thanh lý tài sản, xử lý nợ?', 'VSA 330', 1.0, 'G441', 'Chứng từ thu nhập khác hợp lệ'),
        ]
    },
}

def style_cell(cell, font=None, alignment=None, fill=None, border=None, num_format=None):
    if font: cell.font = font
    if alignment: cell.alignment = alignment
    if fill: cell.fill = fill
    if border: cell.border = border
    if num_format: cell.number_format = num_format

def inject_scoring_sheet_into_file(code, filename, info):
    filepath = os.path.join(TARGET_DIR, filename)
    if not os.path.exists(filepath):
        print(f"[SKIP] File not found: {filepath}")
        return False

    wb = openpyxl.load_workbook(filepath)
    sheet_name = info['sheet_name']

    # Nếu sheet đã tồn tại, xóa đi để tạo mới sạch sẽ
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]

    ws = wb.create_sheet(title=sheet_name)
    doc_font_name = 'Cambria'

    # Styles
    font_main_title = Font(name=doc_font_name, size=13, bold=True, color="002060")
    font_sub_title = Font(name=doc_font_name, size=10, bold=True, color="1E293B")
    font_meta = Font(name=doc_font_name, size=10, bold=False, color="334155")
    font_meta_bold = Font(name=doc_font_name, size=10, bold=True, color="0F172A")
    font_tbl_head = Font(name=doc_font_name, size=10, bold=True, color="000000")
    font_body = Font(name=doc_font_name, size=10, bold=False, color="000000")
    font_body_bold = Font(name=doc_font_name, size=10, bold=True, color="000000")
    font_ref = Font(name=doc_font_name, size=10, bold=True, color="002060")
    font_status_ok = Font(name=doc_font_name, size=10, bold=True, color="059669")
    font_summary_val = Font(name=doc_font_name, size=12, bold=True, color="1D4ED8")

    fill_header = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_summary_box = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    fill_total_row = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="000000")
    double_bottom_side = Side(border_style="double", color="000000")

    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    tbl_bottom_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=double_bottom_side)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # 1. Header thông tin đơn vị & HSKT
    ws.row_dimensions[1].height = 20
    c_firm = ws.cell(1, 1, "CÔNG TY TNHH KIỂM TOÁN BẮC ĐẨU")
    style_cell(c_firm, font_meta_bold, align_left)

    c_ref = ws.cell(1, 7, f"Ký hiệu: {info['ref_code']}")
    style_cell(c_ref, font_meta_bold, align_right)

    ws.row_dimensions[2].height = 18
    c_client = ws.cell(2, 1, "Khách hàng: Công ty Cổ phần / TNHH Khách Hàng")
    style_cell(c_client, font_meta, align_left)

    c_date = ws.cell(2, 7, "Niên độ: 31/12/2025")
    style_cell(c_date, font_meta, align_right)

    ws.row_dimensions[3].height = 18
    c_perf = ws.cell(3, 1, "Người thực hiện: Nguyễn Đắc Thịnh")
    style_cell(c_perf, font_meta, align_left)

    c_rev = ws.cell(3, 7, "Người soát xét: Trưởng nhóm kiểm toán")
    style_cell(c_rev, font_meta, align_right)

    # 2. Tiêu đề Bảng Chấm Điểm
    ws.row_dimensions[5].height = 24
    c_title = ws.cell(5, 1, "BẢNG TỰ CHẤM ĐIỂM & RÀ SOÁT CHẤT LƯỢNG HỒ SƠ KIỂM TOÁN (THEO CHUẨN VACPA)")
    style_cell(c_title, font_main_title, align_left)

    ws.row_dimensions[6].height = 18
    c_subtitle = ws.cell(6, 1, f"Áp dụng cho: {info['section_title']}")
    style_cell(c_subtitle, font_sub_title, align_left)

    # 3. Khung KPI Tóm Tắt Điểm Số (Dòng 8-9)
    ws.row_dimensions[8].height = 20
    ws.row_dimensions[9].height = 24

    ws.merge_cells("A8:B8")
    ws.cell(8, 1, "TỔNG ĐIỂM CHUẨN VACPA")
    style_cell(ws.cell(8, 1), font_meta_bold, align_center, fill_summary_box, cell_border)
    style_cell(ws.cell(8, 2), border=cell_border)

    ws.merge_cells("C8:D8")
    ws.cell(8, 3, "ĐIỂM KTV TỰ CHẤM")
    style_cell(ws.cell(8, 3), font_meta_bold, align_center, fill_summary_box, cell_border)
    style_cell(ws.cell(8, 4), border=cell_border)

    ws.merge_cells("E8:F8")
    ws.cell(8, 5, "TỶ LỆ HOÀN THÀNH")
    style_cell(ws.cell(8, 5), font_meta_bold, align_center, fill_summary_box, cell_border)
    style_cell(ws.cell(8, 6), border=cell_border)

    ws.merge_cells("G8:H8")
    ws.cell(8, 7, "KẾT QUẢ TỰ ĐÁNH GIÁ")
    style_cell(ws.cell(8, 7), font_meta_bold, align_center, fill_summary_box, cell_border)
    style_cell(ws.cell(8, 8), border=cell_border)

    # Giá trị KPI dòng 9
    crit_count = len(info['criteria'])
    start_row = 12
    end_row = start_row + crit_count - 1
    total_row = end_row + 1

    ws.merge_cells("A9:B9")
    ws.cell(9, 1, f"=D{total_row}")
    style_cell(ws.cell(9, 1), font_summary_val, align_center, fill_summary_box, cell_border, "0.0")
    style_cell(ws.cell(9, 2), border=cell_border)

    ws.merge_cells("C9:D9")
    ws.cell(9, 3, f"=F{total_row}")
    style_cell(ws.cell(9, 3), font_summary_val, align_center, fill_summary_box, cell_border, "0.0")
    style_cell(ws.cell(9, 4), border=cell_border)

    ws.merge_cells("E9:F9")
    ws.cell(9, 5, f"=IF(A9>0, C9/A9, 1.0)")
    style_cell(ws.cell(9, 5), font_summary_val, align_center, fill_summary_box, cell_border, "0.0%")
    style_cell(ws.cell(9, 6), border=cell_border)

    ws.merge_cells("G9:H9")
    ws.cell(9, 7, f'=IF(C9>=A9, "ĐẠT YÊU CẦU", "CẦN BỔ SUNG THỦ TỤC")')
    style_cell(ws.cell(9, 7), font_status_ok, align_center, fill_summary_box, cell_border)
    style_cell(ws.cell(9, 8), border=cell_border)

    # 4. Table Header (Dòng 11)
    ws.row_dimensions[11].height = 26
    headers = [
        ("TT", 6, align_center),
        ("NỘI DUNG TIÊU CHÍ ĐÁNH GIÁ (THEO BIỂU CHẤM ĐIỂM BỘ TÀI CHÍNH)", 52, align_left),
        ("CHUẨN MỰC VSA", 15, align_center),
        ("ĐIỂM CHUẨN", 11, align_right),
        ("ĐÃ LÀM?", 12, align_center),
        ("TỰ CHẤM", 11, align_right),
        ("THAM CHIẾU GLV", 16, align_center),
        ("GHI CHÚ & BẰNG CHỨNG BỔ SUNG", 32, align_left)
    ]

    for c_idx, (h_title, col_w, h_align) in enumerate(headers, start=1):
        c = ws.cell(11, c_idx, h_title)
        style_cell(c, font_tbl_head, align_center, fill_header, cell_border)
        col_letter = get_column_letter(c_idx)
        ws.column_dimensions[col_letter].width = col_w

    # 5. Điền từng tiêu chí (Dòng 12 đến end_row)
    for idx, (tt, question, vsa, pts, ref_glv, note) in enumerate(info['criteria'], start=start_row):
        ws.row_dimensions[idx].height = 28

        c_tt = ws.cell(idx, 1, tt)
        style_cell(c_tt, font_body_bold, align_center, border=cell_border)

        c_q = ws.cell(idx, 2, question)
        style_cell(c_q, font_body, align_left, border=cell_border)

        c_vsa = ws.cell(idx, 3, vsa)
        style_cell(c_vsa, font_body, align_center, border=cell_border)

        c_pts = ws.cell(idx, 4, pts)
        style_cell(c_pts, font_body, align_right, border=cell_border, num_format="0.0")

        c_done = ws.cell(idx, 5, "Đã làm")
        style_cell(c_done, font_status_ok, align_center, border=cell_border)

        # Điểm tự chấm = Điểm chuẩn khi KTV đã hoàn thành
        c_self = ws.cell(idx, 6, f"=IF(E{idx}=\"Đã làm\", D{idx}, 0)")
        style_cell(c_self, font_body_bold, align_right, border=cell_border, num_format="0.0")

        c_ref = ws.cell(idx, 7, ref_glv)
        style_cell(c_ref, font_ref, align_center, border=cell_border)

        c_note = ws.cell(idx, 8, note)
        style_cell(c_note, font_body, align_left, border=cell_border)

    # 6. Dòng Tổng Cộng
    ws.row_dimensions[total_row].height = 24

    ws.merge_cells(f"A{total_row}:C{total_row}")
    c_tot_lbl = ws.cell(total_row, 1, f"TỔNG CỘNG ĐIỂM {info['ref_code']} (TỐI ĐA {info['max_points']:.1f} ĐIỂM)")
    style_cell(c_tot_lbl, font_body_bold, align_center, fill_total_row, tbl_bottom_border)
    style_cell(ws.cell(total_row, 2), border=tbl_bottom_border)
    style_cell(ws.cell(total_row, 3), border=tbl_bottom_border)

    c_tot_max = ws.cell(total_row, 4, f"=SUM(D{start_row}:D{end_row})")
    style_cell(c_tot_max, font_body_bold, align_right, fill_total_row, tbl_bottom_border, num_format="0.0")

    c_tot_status = ws.cell(total_row, 5, "")
    style_cell(c_tot_status, font_body_bold, align_center, fill_total_row, tbl_bottom_border)

    # Ô TỔNG ĐIỂM TỰ CHẤM (Ô NÀY SẼ ĐƯỢC LINK SANG SHEET H 110 FILE A-B-H)
    c_tot_self = ws.cell(total_row, 6, f"=SUM(F{start_row}:F{end_row})")
    style_cell(c_tot_self, font_summary_val, align_right, fill_total_row, tbl_bottom_border, num_format="0.0")

    ws.merge_cells(f"G{total_row}:H{total_row}")
    c_tot_eval = ws.cell(total_row, 7, f'=IF(F{total_row}>=D{total_row}, "XẾP LOẠI: XUẤT SẮC", "CẦN BỔ SUNG")')
    style_cell(c_tot_eval, font_status_ok, align_center, fill_total_row, tbl_bottom_border)
    style_cell(ws.cell(total_row, 8), border=tbl_bottom_border)

    # 7. Khung Ký Tên Nghiệm Thu (Cuối Sheet)
    sign_r1 = total_row + 3
    sign_r2 = sign_r1 + 1
    sign_r3 = sign_r1 + 4

    ws.cell(sign_r1, 2, "KIỂM TOÁN VIÊN THỰC HIỆN")
    style_cell(ws.cell(sign_r1, 2), font_body_bold, align_center)
    ws.cell(sign_r2, 2, "(Ký và ghi rõ họ tên)")
    style_cell(ws.cell(sign_r2, 2), font_meta, align_center)
    ws.cell(sign_r3, 2, "Nguyễn Đắc Thịnh")
    style_cell(ws.cell(sign_r3, 2), font_body_bold, align_center)

    ws.cell(sign_r1, 7, "TRƯỞNG NHÓM SOÁT XÉT CHẤT LƯỢNG")
    style_cell(ws.cell(sign_r1, 7), font_body_bold, align_center)
    ws.cell(sign_r2, 7, "(Ký và ghi rõ họ tên)")
    style_cell(ws.cell(sign_r2, 7), font_meta, align_center)
    ws.cell(sign_r3, 7, "Phụ trách cuộc kiểm toán")
    style_cell(ws.cell(sign_r3, 7), font_body_bold, align_center)

    wb.save(filepath)
    print(f"[SUCCESS] Created sheet [{sheet_name}] in {filename} (Total: {info['max_points']:.1f} pts at F{total_row})")
    return (code, sheet_name, f"F{total_row}")

def main():
    print("=== BẮT ĐẦU TẠO SHEET TỰ CHẤM ĐIỂM CHUẨN VACPA VÀO 13 FILE GLV ===")
    files_map = [
        ('D100', 'D100 - Tien - Mau 2024 - Thinh.xlsx'),
        ('D200', 'D200 - Dau tu - ABC 2020.xlsx'),
        ('D300', 'D300 - Phai thu - Mau 2025 - Thinh.xlsx'),
        ('D500', 'D500 - HTK - Mau 2024 - Thinh.xlsx'),
        ('D600', 'D600 - Phan bo - Mau 2024 - Thinh.xlsx'),
        ('D700', 'D700 - Tai san - Mau 2024 - Thinh.xlsx'),
        ('E100', 'E100 - Vay - Mau 2024 - Thinh.xlsx'),
        ('E200', 'E200 - Phai tra - Mau 2024 - Thinh.xlsx'),
        ('E300', 'E300 - Thue - Mau 2024 - Thinh.xlsx'),
        ('E400', 'E400 - Luong - Mau 2025 - Thinh.xlsx'),
        ('F100', 'F100 - Von - Mau 2024 - Thinh.xlsx'),
        ('G100', 'G100 - Doanh thu - Mau 2025- Thinh.xlsx'),
        ('G200', 'G200 - 300 - 400 -  Mau 2025 - Thinh.xlsx'),
    ]

    results = []
    for code, filename in files_map:
        info = SCORING_SHEETS_DATABASE.get(code)
        if not info:
            continue
        res = inject_scoring_sheet_into_file(code, filename, info)
        if res:
            results.append(res)

    print(f"\n=== HOÀN TẤT: Đã tạo thành công {len(results)}/13 sheet tự chấm điểm ===")
    print("\n--- BẢNG TỌA ĐỘ LINK Ô ĐIỂM SANG FILE A - B - H (SHEET H 110) ---")
    for code, sname, cell_coord in results:
        print(f"  Phần hành {code}: ='{sname}'!{cell_coord}")

if __name__ == '__main__':
    main()
